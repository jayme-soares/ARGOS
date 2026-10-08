import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from argos.planilha import converter_data, ler_exportacao_campo, ler_religas_em_campo

CABECALHO = [
    "CO", "Código TdC", "Numero de Serviço", " Tipo de Serviço", "Código Cliente",
    "Município", "Bairro", "Código Equipe", "Prazo ANS Legal", "Endereço Completo",
]
LINHAS = [
    ["132", 111222333, 9000001, "RELIGACAO URGENTE", 55501.0, "MARICA", "Centro", "NI201LRP-B", "29/09/2026 14:30", "Rua A, 1"],
    ["132", 111222334, 9000002, "RELIGACAO NORMAL", 55502, "Maricá", "Itaipuaçu", "NI202LRP-B", datetime(2026, 9, 29, 11, 5), "Rua B, 2"],
    ["132", 111222334, 9000002, "RELIGACAO NORMAL", 55502, "MARICA", "Itaipuaçu", "NI202LRP-B", "29/09/2026 11:05", "duplicada"],
    ["132", 111222335, 9000003, "RELIGACAO NORMAL", 55503, "MARICA", "Inoã", "ni215LRP-B", None, ""],
    # fora do escopo: equipe de outra empresa, sem equipe, outro município
    ["132", 111222336, 9000004, "RELIGACAO NORMAL", 55504, "MARICA", "Centro", "084941", "29/09/2026 09:00", ""],
    ["132", 111222337, 9000005, "RELIGACAO NORMAL", 55505, "MARICA", "Centro", "", "29/09/2026 09:00", ""],
    ["132", 111222338, 9000006, "RELIGACAO NORMAL", 55506, "NITEROI", "Icaraí", "NI201LRP-B", "29/09/2026 09:00", ""],
]


class TestPlanilha(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def conferir(self, registros):
        self.assertEqual([r["tdc"] for r in registros], ["111222334", "111222333", "111222335"])
        primeiro = registros[0]
        self.assertEqual(primeiro["ordem"], "9000002")
        self.assertEqual(primeiro["cliente"], "55502")
        self.assertEqual(primeiro["tipo"], "RELIGACAO NORMAL")
        self.assertEqual(primeiro["equipe"], "NI202LRP-B")
        self.assertEqual(primeiro["bairro"], "Itaipuaçu")
        self.assertTrue(primeiro["vencimento"].startswith("2026-09-29T11:05"))
        self.assertEqual(registros[1]["cliente"], "55501")
        self.assertIsNone(registros[2]["vencimento"])

    def test_xlsx_com_titulo_antes_do_cabecalho(self):
        wb = Workbook()
        wb.active.title = "Filtros"
        wb.active.append(["FILTROS"])
        ws = wb.create_sheet("TdC")
        ws.append(["Relatório exportado pelo eOrder"])
        ws.append([])
        ws.append(CABECALHO)
        for linha in LINHAS:
            ws.append(linha)
        caminho = self.dir / "export.xls"  # extensão enganosa de propósito
        wb.save(caminho)
        self.conferir(ler_religas_em_campo(caminho))

    def test_html_disfarcado_de_xls(self):
        def td(v):
            if isinstance(v, datetime):
                v = v.strftime("%d/%m/%Y %H:%M")
            return f"<td>{'' if v is None else v}</td>"
        html = "<html><body><table><tr>" + "".join(f"<th>{c}</th>" for c in CABECALHO) + "</tr>"
        for linha in LINHAS:
            html += "<tr>" + "".join(td(v) for v in linha) + "</tr>"
        html += "</table></body></html>"
        caminho = self.dir / "export.xls"
        caminho.write_text(html, encoding="utf-8")
        self.conferir(ler_religas_em_campo(caminho))

    def test_separa_finalizadas(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "TdC"
        ws.append(["Código TdC", "Numero de Serviço", "Tipo de Serviço", "Código Cliente", "Município",
                   "Bairro", "Código Equipe", "Prazo ANS Legal", "Estado TdC"])
        ws.append([1001, "A1", "OUTREA", 1, "MARICA", "Centro", "NI201LRP-B", "30/09/2026 12:00", "Enviado ao Campo"])
        ws.append([1002, "A2", "OUTREA", 2, "MARICA", "Centro", "NI202LRP-B", "30/09/2026 12:00", "Finalizado"])
        ws.append([1003, "A3", "OUTREA", 3, "MARICA", "Inoã", "NI203LRP-B", "30/09/2026 12:00", " ENCERRADO "])
        ws.append([1004, "A4", "OUTREA", 4, "MARICA", "Inoã", "NI204LRP-B", "30/09/2026 12:00", "Concluído"])
        # fora do escopo: finalizadas de outro município e de equipe de outra empresa
        ws.append([1005, "A5", "OUTREA", 5, "NITEROI", "Icaraí", "NI201LRP-B", "30/09/2026 12:00", "Finalizado"])
        ws.append([1006, "A6", "OUTREA", 6, "MARICA", "Centro", "084941", "30/09/2026 12:00", "Finalizado"])
        linhas = wb.create_sheet("Linhas TdC")
        linhas.append(["Código TdC", "Resultado", "Causa/Descritivo Resultado", "Data Fim"])
        linhas.append([1002, "Não Realizado", "FJL - Fim da Jornada Laborativa", "29/09/2026 17:00"])
        linhas.append([1002, "Realizado", "", "30/09/2026 10:15"])  # a mais recente vale
        linhas.append([1003, "Realizado", "", "30/09/2026 13:40"])  # depois do prazo
        caminho = self.dir / "export.xlsx"
        wb.save(caminho)

        dados = ler_exportacao_campo(caminho)
        self.assertEqual([r["tdc"] for r in dados["em_aberto"]], ["1001"])
        finalizadas = {r["tdc"]: r for r in dados["finalizadas"]}
        self.assertEqual(list(finalizadas), ["1003", "1002", "1004"])  # mais recente primeiro, sem data no fim
        self.assertTrue(finalizadas["1002"]["finalizada_em"].startswith("2026-09-30T10:15"))
        self.assertEqual(finalizadas["1002"]["resultado"], "Realizado")
        self.assertTrue(finalizadas["1002"]["no_prazo"])
        self.assertFalse(finalizadas["1003"]["no_prazo"])
        self.assertIsNone(finalizadas["1004"]["finalizada_em"])
        self.assertIsNone(finalizadas["1004"]["no_prazo"])
        self.assertEqual(dados["estados"]["Finalizado"], 3)
        self.assertEqual(dados["estados"]["Enviado ao Campo"], 1)
        self.assertEqual([r["tdc"] for r in ler_religas_em_campo(caminho)], ["1001"])

    def test_coluna_faltando_explica(self):
        wb = Workbook()
        wb.active.append(["Código TdC", "Código Equipe"])
        wb.active.append([1, "EQ"])
        caminho = self.dir / "x.xlsx"
        wb.save(caminho)
        with self.assertRaises(ValueError) as ctx:
            ler_religas_em_campo(caminho)
        self.assertIn("Prazo ANS Legal", str(ctx.exception))

    def test_converter_data(self):
        self.assertEqual(converter_data("29/09/2026 14:30").hour, 14)
        self.assertEqual(converter_data("29/09/2026 14:30:15").minute, 30)
        self.assertEqual(converter_data(46294.5).hour, 12)  # serial do Excel
        self.assertIsNone(converter_data(""))
        self.assertIsNone(converter_data(float("nan")))


if __name__ == "__main__":
    unittest.main()
