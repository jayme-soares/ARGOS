import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from argos import composicao
from argos.contas_equipes import codigo_valido, email_da_equipe
from argos.contratos import NEGOCIACAO
from argos.planilha import ler_exportacao_campo
from tests.test_contratos import CABECALHO

PESSOAS = [
    (21687, "CARLOS ALBERTO MENDONCA", "SUL"),
    (11778, "PEDRO HENRIQUE DOS SANTOS", "PETROPOLIS"),
    (11813, "HÉLIO DE SOUZA", "SUL"),
    (14779, "EVERSON GOMES LESSA", "PETROPOLIS"),
    (20001, "JOAO SILVA", "SUL"),
    (20002, "JOAO PEDRO SILVA", "SUL"),
    (20003, "MARCUS SILVA COSTA", "SUL"),
    (20004, "MARCUS SILVA RAMOS", "SUL"),
]


class TestComposicao(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        wb = Workbook()
        wb.active.title = "Planilha1"
        ws = wb.create_sheet("Base")
        ws.append(["Matrícula", "Equipe", "Encarregado", "Município"])
        for m, nome, mun in PESSOAS:
            ws.append([m, nome, "X", mun])
        self.xlsx = self.dir / "composicao.xlsx"
        wb.save(self.xlsx)
        self.cadastro = composicao.ler_planilha(self.xlsx)

    def tearDown(self):
        self.tmp.cleanup()

    def test_le_aba_base(self):
        self.assertEqual(len(self.cadastro), len(PESSOAS))
        self.assertEqual(self.cadastro["11813"], {"nome": "HÉLIO DE SOUZA", "municipio": "SUL"})

    def test_nome_do_eorder_abreviado_e_fora_de_ordem(self):
        m = lambda nome: composicao.matricula_do_nome(nome, self.cadastro)
        self.assertEqual(m("ALBERTO MENDONÇA CARLOS"), "21687")
        self.assertEqual(m("PEDRO SANTOS"), "11778")
        self.assertEqual(m("HÉLIO  DE SOUZA"), "11813")
        self.assertEqual(m("GOMES LESSA EVERSON"), "14779")
        # dois candidatos: vale o de nome idêntico; sem idêntico, ninguém
        self.assertEqual(m("JOAO SILVA"), "20001")
        self.assertIsNone(m("MARCUS SILVA"))
        self.assertIsNone(m("FULANO DE TAL"))
        self.assertIsNone(m(""))

    def test_matricula(self):
        self.assertEqual(composicao.normalizar_matricula(20045.0), "20045")
        self.assertEqual(composicao.normalizar_matricula(" 020045 "), "20045")
        self.assertEqual(composicao.normalizar_matricula(None), "")
        self.assertTrue(codigo_valido("20045", NEGOCIACAO))
        self.assertFalse(codigo_valido("CENEGED-ANGRA-01", NEGOCIACAO))
        self.assertEqual(email_da_equipe("20045", NEGOCIACAO), "neg-20045@equipes.argos.local")

    def test_bot_usa_a_matricula_como_equipe(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "TdC"
        ws.append(CABECALHO)
        ws.append([1, "A1", "RELIGA", 1, "ANGRA DOS REIS", "Centro", "083446", "ALBERTO MENDONÇA CARLOS",
                   "30/09/2026 12:00", "Enviado ao Campo"])
        ws.append([2, "A2", "RELIGA", 2, "ANGRA DOS REIS", "Centro", "079491", "RONDINELI SILVA",
                   "30/09/2026 13:00", "Enviado ao Campo"])
        exportacao = self.dir / "neg.xlsx"
        wb.save(exportacao)
        with mock.patch.dict(os.environ, {"ARGOS_NEG_COMPOSICAO": str(self.xlsx)}):
            regs = ler_exportacao_campo(exportacao, NEGOCIACAO)["em_aberto"]
        self.assertEqual([(r["equipe"], r["equipe_nome"]) for r in regs],
                         [("21687", "ALBERTO MENDONÇA CARLOS"), ("RONDINELI-SILVA", "RONDINELI SILVA")])


if __name__ == "__main__":
    unittest.main()
