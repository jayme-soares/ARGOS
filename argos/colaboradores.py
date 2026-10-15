"""Colaboradores da Negociação (planilha "Colaboradores.xlsx").

Colunas: REGIONAL, MATRÍCULA, Codigo Equipe, ACESSO EORDER e REFUN (nome
completo). A matrícula é o código de login da equipe no painel /equipe e a
chave da equipe nos avisos. O "Codigo Equipe" é o mesmo "Código Equipe" da
exportação do eOrder: é por ele que cada ordem chega à matrícula.

A planilha tem alguns desvios, tratados assim:
- código com sufixo ("079471_653"): vale a parte antes do "_";
- linha repetida (mesma matrícula e código): conta uma vez;
- código em mais de uma matrícula: desempata pelo nome do eOrder (todas as
  palavras no nome completo; depois, igual ao ACESSO EORDER); se ainda
  sobrar mais de um, a ordem fica sem matrícula;
- sem código: a conta existe, mas nenhuma ordem chega a ela.

O bot lê a planilha em DATA_DIR/negociacao/colaboradores.xlsx (ou
ARGOS_NEG_COLABORADORES) e relê quando o arquivo muda. Ordem sem matrícula
fica com o nome em forma de slug como chave, como antes.
"""

import os
import re
import threading
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from argos import config

ARQUIVO = "negociacao/colaboradores.xlsx"
CONECTIVOS = {"DA", "DE", "DO", "DAS", "DOS", "E"}

_lock = threading.Lock()
_cache: dict = {"chave": None, "cadastro": None}


@dataclass(frozen=True)
class Colaborador:
    matricula: str
    nome: str           # REFUN (nome completo)
    nome_eorder: str    # ACESSO EORDER
    codigo: str         # Codigo Equipe normalizado ("" se não tem)
    regional: str


def caminho_padrao() -> Path:
    return Path(os.environ.get("ARGOS_NEG_COLABORADORES") or config.DATA_DIR / ARQUIVO)


def _digitos(valor) -> str:
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    d = re.sub(r"\D", "", str(valor if valor is not None else ""))
    return d.lstrip("0") or ("0" if d else "")


def normalizar_matricula(valor) -> str:
    """20045, "20045", "020045", 20045.0 -> "20045". Igual a
    normalizarMatricula() em web/api/_equipes.js."""
    return _digitos(valor)


def normalizar_codigo(valor) -> str:
    """Código Equipe sem zeros à esquerda e sem sufixo: "079471_653" -> "79471"."""
    return _digitos(str(valor if valor is not None else "").split("_")[0])


def _sem_acento(texto) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", str(texto or "")) if not unicodedata.combining(c)).upper()


def palavras(nome) -> frozenset[str]:
    return frozenset(re.findall(r"[A-Z0-9]+", _sem_acento(nome))) - CONECTIVOS


def ler_planilha(caminho: Path | str) -> list[Colaborador]:
    """Um Colaborador por (matrícula, código), na ordem da planilha."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        linhas = ws.iter_rows(values_only=True)
        cab = [_sem_acento(c).strip() for c in next(linhas, ())]

        def col(*nomes):
            return next((i for i, c in enumerate(cab) if c in nomes), None)

        i_mat, i_cod, i_eorder, i_nome, i_reg = (col("MATRICULA"), col("CODIGO EQUIPE"), col("ACESSO EORDER"),
                                                 col("REFUN", "NOME"), col("REGIONAL"))
        if i_mat is None or i_cod is None:
            raise ValueError(f"faltam as colunas MATRÍCULA e Codigo Equipe (cabeçalho: {cab})")
        pega = lambda l, i: l[i] if i is not None and i < len(l) else None
        vistos, saida = set(), []
        for l in linhas:
            mat = normalizar_matricula(pega(l, i_mat))
            if not mat:
                continue
            cod = normalizar_codigo(pega(l, i_cod))
            if (mat, cod) in vistos:
                continue
            vistos.add((mat, cod))
            nome_eorder = " ".join(str(pega(l, i_eorder) or "").split())
            saida.append(Colaborador(mat, " ".join(str(pega(l, i_nome) or "").split()) or nome_eorder,
                                     nome_eorder, cod, str(pega(l, i_reg) or "").strip()))
        return saida
    finally:
        wb.close()


def cadastro(caminho: Path | str | None = None) -> list[Colaborador]:
    """Planilha atual, relida só quando o arquivo muda. [] se não existe."""
    arq = Path(caminho) if caminho else caminho_padrao()
    try:
        st = arq.stat()
    except OSError:
        return []
    chave = (str(arq), st.st_mtime_ns, st.st_size)
    with _lock:
        if _cache["chave"] != chave:
            try:
                _cache["cadastro"] = ler_planilha(arq)
            except Exception as e:
                print(f"[colaboradores] não consegui ler {arq}: {e}", flush=True)
                _cache["cadastro"] = []
            _cache["chave"] = chave
        return _cache["cadastro"]


def candidatos(codigo_equipe, colaboradores: list[Colaborador]) -> list[Colaborador]:
    cod = normalizar_codigo(codigo_equipe)
    return [c for c in colaboradores if cod and c.codigo == cod]


def matricula(codigo_equipe, nome_eorder=None, colaboradores: list[Colaborador] | None = None) -> str | None:
    """Matrícula do "Código Equipe" do eOrder (ver o desempate no topo)."""
    colaboradores = cadastro() if colaboradores is None else colaboradores
    achados = {c.matricula: c for c in candidatos(codigo_equipe, colaboradores)}
    alvo = palavras(nome_eorder)
    # Desempate: primeiro o nome completo (o ACESSO EORDER às vezes foi
    # copiado de outra linha), depois o ACESSO EORDER.
    for bate in (lambda c: alvo <= palavras(c.nome), lambda c: palavras(c.nome_eorder) == alvo):
        if len(achados) <= 1 or not alvo:
            break
        filtrados = {m: c for m, c in achados.items() if bate(c)}
        achados = filtrados or achados
    return next(iter(achados)) if len(achados) == 1 else None
