"""Composição das equipes da Negociação (planilha "Composição Fixa").

A aba "Base" tem a Matrícula e o nome completo de cada eletricista. A
matrícula é o código de login da equipe no painel /equipe e a chave da
equipe nos avisos. O eOrder não tem a matrícula: a coluna "Equipe" traz o
nome, muitas vezes abreviado e com a ordem trocada ("ALBERTO MENDONÇA
CARLOS" = CARLOS ALBERTO MENDONCA, "PEDRO SANTOS" = PEDRO HENRIQUE DOS
SANTOS). Por isso o nome do eOrder casa com quem tem TODAS as palavras dele
no nome completo, e só quando há um único candidato.

O bot lê a planilha em DATA_DIR/negociacao/composicao.xlsx (ou
ARGOS_NEG_COMPOSICAO) e relê quando o arquivo muda. Sem a planilha, ou sem
casamento, a equipe fica com o nome em forma de slug, como antes.
"""

import os
import re
import threading
import unicodedata
from pathlib import Path

from argos import config

ARQUIVO = "negociacao/composicao.xlsx"
CONECTIVOS = {"DA", "DE", "DO", "DAS", "DOS", "E"}

_lock = threading.Lock()
_cache: dict = {"chave": None, "pessoas": {}}


def caminho_padrao() -> Path:
    return Path(os.environ.get("ARGOS_NEG_COMPOSICAO") or config.DATA_DIR / ARQUIVO)


def normalizar_matricula(valor) -> str:
    """20045, "20045", "020045", 20045.0 -> "20045". Igual a
    normalizarCodigo() da Negociação em web/api/_equipes.js."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    digitos = re.sub(r"\D", "", str(valor if valor is not None else ""))
    return digitos.lstrip("0") or ("0" if digitos else "")


def _sem_acento(texto) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", str(texto or "")) if not unicodedata.combining(c)).upper()


def palavras(nome) -> frozenset[str]:
    return frozenset(re.findall(r"[A-Z0-9]+", _sem_acento(nome))) - CONECTIVOS


def ler_planilha(caminho: Path | str) -> dict[str, dict]:
    """{matrícula: {"nome", "municipio"}} da aba "Base" (cabeçalho na 1ª linha
    com "Matrícula" e "Equipe")."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb["Base"] if "Base" in wb.sheetnames else wb.worksheets[0]
        linhas = ws.iter_rows(values_only=True)
        cab = [_sem_acento(c).strip() for c in next(linhas, ())]
        i_mat = next(i for i, c in enumerate(cab) if "MATRICULA" in c)
        i_nome = next(i for i, c in enumerate(cab) if c == "EQUIPE")
        i_mun = next((i for i, c in enumerate(cab) if "MUNICIPIO" in c), None)
        pessoas = {}
        for l in linhas:
            mat = normalizar_matricula(l[i_mat] if i_mat < len(l) else None)
            nome = " ".join(str(l[i_nome] or "").split()) if i_nome < len(l) else ""
            if mat and nome:
                pessoas[mat] = {"nome": nome, "municipio": str(l[i_mun] or "").strip() if i_mun is not None else ""}
        return pessoas
    finally:
        wb.close()


def pessoas(caminho: Path | str | None = None) -> dict[str, dict]:
    """Planilha atual, relida só quando o arquivo muda. {} se não existe."""
    arq = Path(caminho) if caminho else caminho_padrao()
    try:
        st = arq.stat()
    except OSError:
        return {}
    chave = (str(arq), st.st_mtime_ns, st.st_size)
    with _lock:
        if _cache["chave"] != chave:
            try:
                _cache["pessoas"] = ler_planilha(arq)
            except Exception as e:
                print(f"[composicao] não consegui ler {arq}: {e}", flush=True)
                _cache["pessoas"] = {}
            _cache["chave"] = chave
        return _cache["pessoas"]


def matricula_do_nome(nome_eorder, cadastro: dict[str, dict] | None = None) -> str | None:
    """Matrícula de quem tem todas as palavras de `nome_eorder` no nome
    completo. Com mais de um candidato, só vale o de nome idêntico."""
    cadastro = pessoas() if cadastro is None else cadastro
    alvo = palavras(nome_eorder)
    if not alvo or not cadastro:
        return None
    candidatos = [m for m, p in cadastro.items() if alvo <= palavras(p["nome"])]
    if len(candidatos) > 1:
        candidatos = [m for m in candidatos if palavras(cadastro[m]["nome"]) == alvo]
    return candidatos[0] if len(candidatos) == 1 else None
